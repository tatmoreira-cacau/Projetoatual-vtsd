from __future__ import annotations

import json
import re
from pathlib import Path

import pdfplumber
from docx import Document
from html import unescape
from PIL import Image


ROOT = Path(__file__).resolve().parent
SKIP_NAMES = {Path(__file__).name, "_project_corpus.json"}


def clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def read_docx(path: Path) -> str:
    doc = Document(path)
    blocks: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            blocks.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return clean("\n".join(blocks))


def read_pdf(path: Path) -> tuple[str, int]:
    pages: list[str] = []
    with pdfplumber.open(path) as pdf:
        for index, page in enumerate(pdf.pages, 1):
            pages.append(f"[PAGE {index}]\n{page.extract_text() or ''}")
        return clean("\n\n".join(pages)), len(pdf.pages)


def read_html(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="replace")
    raw = re.sub(r"<(script|style)\b[^>]*>.*?</\1>", "", raw, flags=re.I | re.S)
    raw = re.sub(r"<br\s*/?>|</(?:p|div|h[1-6]|li|tr)>", "\n", raw, flags=re.I)
    return clean(unescape(re.sub(r"<[^>]+>", " ", raw)))


records = []
for path in sorted(ROOT.rglob("*")):
    if not path.is_file() or path.name in SKIP_NAMES or ".git" in path.parts:
        continue
    rel = path.relative_to(ROOT).as_posix()
    ext = path.suffix.lower()
    record = {"path": rel, "extension": ext, "bytes": path.stat().st_size}
    try:
        if ext in {".md", ".txt"}:
            record["text"] = clean(path.read_text(encoding="utf-8", errors="replace"))
        elif ext == ".html":
            record["text"] = read_html(path)
        elif ext == ".docx":
            record["text"] = read_docx(path)
        elif ext == ".pdf":
            record["text"], record["pages"] = read_pdf(path)
        elif ext in {".jpg", ".jpeg", ".png", ".webp"}:
            with Image.open(path) as image:
                record["image"] = {
                    "width": image.width,
                    "height": image.height,
                    "mode": image.mode,
                    "format": image.format,
                }
    except Exception as exc:
        record["error"] = f"{type(exc).__name__}: {exc}"
    records.append(record)

(ROOT / "_project_corpus.json").write_text(
    json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8"
)
print(json.dumps({
    "files": len(records),
    "text_files": sum("text" in item for item in records),
    "image_files": sum("image" in item for item in records),
    "errors": [item for item in records if "error" in item],
    "characters": sum(len(item.get("text", "")) for item in records),
}, ensure_ascii=False, indent=2))
